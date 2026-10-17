"""
Excel Exporter for Standardized Audit Workpapers (.xlsx).
Enhanced with Execution Mode badge and Three-Layer Evidence Traceability (Issue 5).
"""

from pathlib import Path
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Optional, Dict, Any
from src.core.schemas import AnalysisReportResult, AuditWorkpaper, ExecutionMode


def _cell_text(value: object) -> str:
    """Neutralise Excel formula injection: prefix risky leading chars with an apostrophe."""
    s = "" if value is None else str(value)
    return "'" + s if s[:1] in ("=", "+", "-", "@", "\t", "\r") else s


def export_workpaper_to_excel(report: AnalysisReportResult, output_path: Path, cpa_signoff: Optional[Dict[str, Any]] = None) -> Path:
    """Export Audit Workpapers to a beautifully styled Excel workbook with evidence trail."""
    wb = openpyxl.Workbook()
    
    # Sheet 1: Summary Dashboard
    ws_summary = wb.active
    ws_summary.title = "审计总览与风险矩阵"
    ws_summary.views.sheetView[0].showGridLines = True

    # Styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="微软雅黑", size=16, bold=True, color="1F497D")
    bold_font = Font(name="微软雅黑", size=10, bold=True)
    normal_font = Font(name="微软雅黑", size=10)
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Title
    ws_summary["A1"] = _cell_text(f"{report.company_name} - 数智审计穿透与风险研判底稿")
    ws_summary["A1"].font = title_font
    ws_summary.merge_cells("A1:H1")

    # Meta & Execution Mode
    mode_text = f"🟢 实时在线推理 ({report.model_name})" if report.execution_mode in (ExecutionMode.ONLINE, ExecutionMode.STRICT_ONLINE) and not report.fallback_occurred else "🟡 离线确定性模拟 (Demo Mode)"
    ws_summary["A3"] = "被审计单位:"
    ws_summary["B3"] = _cell_text(report.company_name)
    ws_summary["C3"] = "综合风险评级:"
    ws_summary["D3"] = _cell_text(report.overall_risk_rating.value)
    ws_summary["E3"] = "Beneish M-Score:"
    ws_summary["F3"] = str(report.beneish_m_score if report.beneish_m_score is not None else "N/A")
    ws_summary["G3"] = "执行模式:"
    ws_summary["H3"] = mode_text
    for cell in ["A3", "C3", "E3", "G3"]:
        ws_summary[cell].font = bold_font

    # Executive Summary
    ws_summary["A5"] = "管理层与主审评委摘要:"
    ws_summary["A5"].font = bold_font
    ws_summary["A6"] = _cell_text(report.executive_summary)
    ws_summary["A6"].font = normal_font
    ws_summary.merge_cells("A6:H7")
    ws_summary["A6"].alignment = Alignment(wrap_text=True, vertical="top")

    # Risk Findings Table with Evidence Traceability
    ws_summary["A9"] = "重点风险发现清单与三层证据链 (Risk Findings & Evidence Trail)"
    ws_summary["A9"].font = Font(name="微软雅黑", size=12, bold=True, color="C00000")

    headers_finding = ["编号", "风险发现项", "风险等级", "涉案金额(元)", "确定性客观事实(规则发现)", "大模型准则深度研判", "待人工复核/审计程序", "依据准则"]
    for col_idx, h in enumerate(headers_finding, start=1):
        c = ws_summary.cell(row=10, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 11
    for f in report.findings:
        ws_summary.cell(row=current_row, column=1, value=_cell_text(f.finding_id)).font = normal_font
        ws_summary.cell(row=current_row, column=2, value=_cell_text(f.title)).font = bold_font
        ws_summary.cell(row=current_row, column=3, value=_cell_text(f.risk_level.value)).font = bold_font
        ws_summary.cell(row=current_row, column=4, value=f.impact_amount).font = normal_font
        ws_summary.cell(row=current_row, column=5, value=_cell_text(f.rule_evidence or f.suspected_mechanism)).font = normal_font
        ws_summary.cell(row=current_row, column=6, value=_cell_text(f.model_explanation or f.suspected_mechanism)).font = normal_font
        ws_summary.cell(row=current_row, column=7, value=_cell_text(f.human_verification_flag or f.suggested_procedure)).font = normal_font
        ws_summary.cell(row=current_row, column=8, value=_cell_text(f.accounting_standard)).font = normal_font
        
        for col_idx in range(1, 9):
            ws_summary.cell(row=current_row, column=col_idx).border = border_thin
            ws_summary.cell(row=current_row, column=col_idx).alignment = Alignment(wrap_text=True, vertical="top")
        current_row += 1

    # Individual Workpaper Sheets
    for wp_idx, wp in enumerate(report.workpapers, start=1):
        sheet_name = re.sub(r"[\[\]:*?/\\]", "-", f"底稿_{wp.workpaper_id}")[:31] or "底稿"
        ws_wp = wb.create_sheet(title=sheet_name)
        ws_wp.views.sheetView[0].showGridLines = True

        ws_wp["A1"] = _cell_text(wp.title)
        ws_wp["A1"].font = title_font
        ws_wp.merge_cells("A1:H1")

        ws_wp["A3"] = "底稿编号:"
        ws_wp["B3"] = _cell_text(wp.workpaper_id)
        ws_wp["C3"] = "编制人:"
        ws_wp["D3"] = _cell_text(wp.prepared_by)
        ws_wp["E3"] = "核查日期:"
        ws_wp["F3"] = _cell_text(wp.review_date)

        headers_wp = ["凭证号", "记账日期", "业务摘要", "账面金额(元)", "核实确认金额(元)", "差异金额(元)", "审计结论", "数据来源与追溯"]
        for col_idx, h in enumerate(headers_wp, start=1):
            c = ws_wp.cell(row=5, column=col_idx, value=h)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        r_idx = 6
        for r in wp.rows:
            ws_wp.cell(row=r_idx, column=1, value=_cell_text(r.voucher_no)).font = normal_font
            ws_wp.cell(row=r_idx, column=2, value=_cell_text(r.date)).font = normal_font
            ws_wp.cell(row=r_idx, column=3, value=_cell_text(r.summary)).font = normal_font
            ws_wp.cell(row=r_idx, column=4, value=r.ledger_amount).font = normal_font
            ws_wp.cell(row=r_idx, column=5, value=r.verified_amount).font = normal_font
            ws_wp.cell(row=r_idx, column=6, value=r.discrepancy).font = normal_font
            ws_wp.cell(row=r_idx, column=7, value=_cell_text(r.audit_conclusion)).font = normal_font
            ws_wp.cell(row=r_idx, column=8, value=_cell_text(r.evidence_trace or r.source_file or "标准凭证库")).font = normal_font

            for col_idx in range(1, 9):
                ws_wp.cell(row=r_idx, column=col_idx).border = border_thin
            r_idx += 1

        ws_wp.cell(row=r_idx + 1, column=1, value="底稿综合意见:").font = bold_font
        ws_wp.cell(row=r_idx + 1, column=2, value=_cell_text(wp.audit_opinion_summary)).font = normal_font
        ws_wp.merge_cells(start_row=r_idx + 1, start_column=2, end_row=r_idx + 2, end_column=8)

    # Sheet: Population analytics
    analytics = report.analytics or {}
    je = analytics.get("journal_entry_testing") or {}
    if je.get("applicable"):
        ws_je = wb.create_sheet(title="全量分录测试")
        ws_je["A1"] = "会计分录测试（CSA 1141）· 全量覆盖"
        ws_je["A1"].font = title_font
        ws_je.merge_cells("A1:G1")
        ws_je["A2"] = _cell_text(
            f"总体 {je.get('population')} 张凭证，覆盖率 100%；命中风险标识 {je.get('flagged_count')} 张，高风险（≥5 分）{je.get('high_risk_count')} 张。{je.get('method', '')}"
        )
        ws_je.merge_cells("A2:G2")
        ws_je["A2"].alignment = Alignment(wrap_text=True)
        heads = ["规则编码", "规则", "命中张数", "准则依据"]
        for c, h in enumerate(heads, 1):
            cell = ws_je.cell(row=4, column=c, value=h)
            cell.font, cell.fill = header_font, header_fill
        for r, (code, name) in enumerate((je.get("rules") or {}).items(), 5):
            vals = [code, name, (je.get("rule_hits") or {}).get(code, 0), (je.get("rule_basis") or {}).get(code, "")]
            for c, v in enumerate(vals, 1):
                ws_je.cell(row=r, column=c, value=v if isinstance(v, int) else _cell_text(v)).font = normal_font
        start = 6 + len(je.get("rules") or {})
        heads = ["凭证号", "日期", "金额(元)", "风险分", "命中规则", "来源文件", "行号"]
        for c, h in enumerate(heads, 1):
            cell = ws_je.cell(row=start, column=c, value=h)
            cell.font, cell.fill = header_font, header_fill
        for r, f in enumerate(je.get("flagged") or [], start + 1):
            vals = [f.get("voucher_id"), f.get("date"), f.get("amount"), f.get("score"),
                    "、".join(f.get("rule_names") or []), f.get("source_file"), f.get("row_index")]
            for c, v in enumerate(vals, 1):
                ws_je.cell(row=r, column=c, value=v if isinstance(v, (int, float)) else _cell_text(v)).font = normal_font
        for col, w in zip("ABCDEFG", (16, 34, 16, 10, 48, 24, 8)):
            ws_je.column_dimensions[col].width = w
    rt = analytics.get("fund_round_trips") or {}
    if rt.get("loop_count"):
        ws_rt = wb.create_sheet(title="资金往返识别")
        heads = ["流出流水", "流出对手方", "流出金额", "流入流水", "流入对手方", "流入金额", "间隔天数", "模式"]
        for c, h in enumerate(heads, 1):
            cell = ws_rt.cell(row=1, column=c, value=h)
            cell.font, cell.fill = header_font, header_fill
        for r, l in enumerate(rt["loops"], 2):
            vals = [l["out_id"], l["out_party"], l["out_amount"], l["in_id"], l["in_party"], l["in_amount"], l["days"], l["pattern"]]
            for c, v in enumerate(vals, 1):
                ws_rt.cell(row=r, column=c, value=v if isinstance(v, (int, float)) else _cell_text(v)).font = normal_font

    # Sheet: 拟定审计调整分录（AJE）与审定前后冲击
    adj = report.adjustments or {}
    if adj.get("entries"):
        ws_aje = wb.create_sheet(title="拟定审计调整分录")
        ws_aje["A1"] = "拟定审计调整分录（AJE）· 须经注册会计师与管理层沟通后确定"
        ws_aje["A1"].font = title_font
        ws_aje.merge_cells("A1:F1")
        heads = ["分录号", "对应发现", "科目编码", "科目名称", "借方(元)", "贷方(元)", "说明"]
        for c, h in enumerate(heads, 1):
            cell = ws_aje.cell(row=3, column=c, value=h)
            cell.font, cell.fill = header_font, header_fill
        r = 4
        for a in adj["entries"]:
            for l in a["lines"]:
                vals = [a["aje_id"], a["finding_id"], l["account_code"], l["account_name"], l["debit"] or None, l["credit"] or None, l.get("note") or a["description"]]
                for c, v in enumerate(vals, 1):
                    ws_aje.cell(row=r, column=c, value=v if isinstance(v, (int, float)) or v is None else _cell_text(v)).font = normal_font
                r += 1
            r += 1
        imp = adj.get("impact")
        if imp:
            ws_aje.cell(row=r + 1, column=1, value="审定前后冲击").font = bold_font
            rows = [("营业收入", imp["revenue_before"], imp["revenue_after"]),
                    ("营业成本", imp["cost_before"], imp["cost_after"]),
                    ("净利润（近似）", imp["profit_before"], imp["profit_after"])]
            for c, h in enumerate(["项目", "审定前(元)", "审定后(元)"], 1):
                cell = ws_aje.cell(row=r + 2, column=c, value=h)
                cell.font, cell.fill = header_font, header_fill
            for i, (n, b, a2) in enumerate(rows, r + 3):
                ws_aje.cell(row=i, column=1, value=n).font = normal_font
                ws_aje.cell(row=i, column=2, value=b).font = normal_font
                ws_aje.cell(row=i, column=3, value=a2).font = normal_font
            ws_aje.cell(row=r + 7, column=1, value=_cell_text(imp.get("caveat", ""))).font = normal_font
        for col, w in zip("ABCDEFG", (10, 10, 10, 30, 18, 18, 36)):
            ws_aje.column_dimensions[col].width = w

    # Sheet: CSA 1131 人工智能辅助审计过程记录
    ws_ai = wb.create_sheet(title="CSA1131_AI过程记录")
    ws_ai["A1"] = "人工智能工具使用过程与采纳复核记录（CSA 1131 审计工作底稿）"
    ws_ai["A1"].font = title_font
    ws_ai.merge_cells("A1:F1")
    ver = report.verification or {}
    pa = report.privacy_audit or {}
    mat = (report.substantive or {}).get("materiality") or {}
    pa_breakdown = {k: pa.get(k, 0) for k in ("id_or_uscc", "bank_account", "mobile", "email") if pa.get(k)}
    pa_counts_str = f"；共掩码 {pa.get('total', 0)} 处 {pa_breakdown}" if pa.get("total") else (f"；掩码统计 {pa.get('counts')}" if pa.get("counts") else "")
    meta = [
        ("执行模式", report.execution_mode.value if hasattr(report.execution_mode, "value") else str(report.execution_mode)),
        ("模型", report.model_name),
        ("是否降级", "是：" + (report.fallback_reason or "") if report.fallback_occurred else "否"),
        ("Token 用量", str((report.token_usage or {}).get("total", 0))),
        ("耗时(秒)", str(report.execution_time_seconds)),
        ("确定性算子", "三单勾稽；Beneish M-Score；全量分录测试；资金往返；Benford；重要性；银企核对；背靠背贸易；截止与成本配比；第三方回款与预付；职业判断类算子（费用完整性/减值/预计负债/负债核销/商业实质）"),
        ("重要性水平", f"整体 {mat.get('overall_materiality', 0):,.2f} / 实际执行 {mat.get('performance_materiality', 0):,.2f}（{mat.get('benchmark', '')}）" if mat else "—"),
        ("出境脱敏", (pa.get("note") or "") + pa_counts_str),
        ("证据锚定守卫", f"发现 {ver.get('findings_total', 0)} 项，已锚定 {ver.get('grounded', 0)}，未锚定 {ver.get('ungrounded', 0)}，金额待复核 {ver.get('amount_unverified', 0)}"),
        ("定级校准", f"{len((ver.get('calibration') or {}).get('recalibrated', []))} 项被校准" if ver.get("calibration") else "离线确定性模式，无需校准"),
        ("自我修正", ("已触发：问题 %s → %s，%s" % ((ver.get("reflection") or {}).get("issues_before"), (ver.get("reflection") or {}).get("issues_after"), "采纳修正稿" if (ver.get("reflection") or {}).get("adopted") else "保留初稿")) if ver.get("reflection") else "未触发"),
    ]
    for i, (k, v) in enumerate(meta, 3):
        ws_ai.cell(row=i, column=1, value=k).font = bold_font
        ws_ai.cell(row=i, column=2, value=_cell_text(v)).font = normal_font
        ws_ai.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    start = 4 + len(meta)
    heads = ["发现编号", "标题", "等级", "金额(元)", "来源", "证据锚定", "注册会计师复核意见（采纳/修正采纳/不予采纳）"]
    for c, h in enumerate(heads, 1):
        cell = ws_ai.cell(row=start, column=c, value=h)
        cell.font, cell.fill = header_font, header_fill
    item_dispositions = (cpa_signoff or {}).get("finding_dispositions") or {}
    default_disp = "已复核采纳" if cpa_signoff else "待注册会计师复核"
    for r, f in enumerate(report.findings, start + 1):
        src = "确定性算子" if "确定性" in (f.model_explanation or "") else "大模型"
        disp = item_dispositions.get(f.finding_id, default_disp)
        vals = [f.finding_id, f.title, f.risk_level.value, f.impact_amount, src,
                ("已锚定 " + "、".join(f.cited_refs[:4])) if f.grounded else "未锚定", disp]
        for c, v in enumerate(vals, 1):
            ws_ai.cell(row=r, column=c, value=v if isinstance(v, (int, float)) else _cell_text(v)).font = normal_font
    for col, w in zip("ABCDEFG", (12, 44, 8, 18, 12, 30, 36)):
        ws_ai.column_dimensions[col].width = w

    # Sheet: Ethics & CPA Sign-off
    ws_signoff = wb.create_sheet(title="CPA终审签章与合规声明")
    ws_signoff.views.sheetView[0].showGridLines = True
    ws_signoff["A1"] = "中国注册会计师审计底稿终审签章与独立性声明"
    ws_signoff["A1"].font = title_font
    ws_signoff.merge_cells("A1:G1")

    signoff_data = cpa_signoff or {
        "auditor_name": "未签署",
        "cert_no": "—",
        "audit_firm": "—",
        "sign_timestamp": "—",
        "verification_hash": "—",
        "opinion_type": "待注册会计师复核后确定",
        "professional_skepticism_notes": "本底稿尚未经注册会计师签署。AI 发现项仅为初筛结果，须由注册会计师复核原始证据后形成结论。",
        "ethical_declaration": "AI 智能体仅作为实质性测试初筛辅助工具，舞弊疑点与底稿结论的职业判断与法律责任由签字注册会计师及审计项目组承担。",
    }

    meta_rows = [
        ("签字注册会计师:", signoff_data.get("auditor_name", "未签署"), "执业证书编号:", signoff_data.get("cert_no", "—")),
        ("执业会计师事务所:", signoff_data.get("audit_firm", "—"), "终审签章时间:", signoff_data.get("sign_timestamp", "—")),
        ("初拟审计意见类型:", signoff_data.get("opinion_type", "待注册会计师复核后确定"), "防伪核验哈希:", signoff_data.get("verification_hash", "—")),
    ]

    for idx, (lbl1, val1, lbl2, val2) in enumerate(meta_rows, start=3):
        ws_signoff.cell(row=idx, column=1, value=lbl1).font = bold_font
        ws_signoff.cell(row=idx, column=2, value=_cell_text(val1)).font = normal_font
        ws_signoff.cell(row=idx, column=4, value=lbl2).font = bold_font
        ws_signoff.cell(row=idx, column=5, value=_cell_text(val2)).font = normal_font

    ws_signoff["A7"] = "注册会计师职业怀疑复核意见:"
    ws_signoff["A7"].font = bold_font
    ws_signoff["A8"] = _cell_text(signoff_data.get("professional_skepticism_notes", ""))
    ws_signoff["A8"].font = normal_font
    ws_signoff.merge_cells("A8:G9")
    ws_signoff["A8"].alignment = Alignment(wrap_text=True, vertical="top")

    ws_signoff["A11"] = "伦理合规与责任界定声明:"
    ws_signoff["A11"].font = bold_font
    ws_signoff["A12"] = _cell_text(signoff_data.get("ethical_declaration", ""))
    ws_signoff["A12"].font = normal_font
    ws_signoff.merge_cells("A12:G14")
    ws_signoff["A12"].alignment = Alignment(wrap_text=True, vertical="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
