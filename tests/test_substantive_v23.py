"""v2.3：专项实质性程序、重要性、定级校准、调整分录、上传与导出相关测试。"""
import io
import re
from pathlib import Path

from fastapi.testclient import TestClient

from src.api.server import app
from src.benchmark.test_cases import get_case_by_id
from src.core.harness import AccountingAgentHarness
from src.core.schemas import (
    AccountingCaseData, AccountingVoucher, BankFlowRecord, BusinessContract, JournalEntryLine, RiskFinding, RiskLevel,
)
from src.plugins.audit_fraud_plugin.plugin import AuditFraudPlugin, calibrate_model_findings
from src.plugins.audit_fraud_plugin.substantive import (
    compute_materiality, detect_third_party_and_prepay, reconcile_book_to_bank, revenue_net, verify_cutoff_and_cost,
)

client = TestClient(app)


def _v(vid, d, lines, doc=None):
    return AccountingVoucher(voucher_id=vid, voucher_date=d, associated_doc_id=doc,
                             entries=[JournalEntryLine(account_code=c, account_name=n, debit=dr, credit=cr) for c, n, dr, cr in lines])


def _case(vouchers, flows=(), contracts=()):
    return AccountingCaseData(case_id="T-1", company_name="测试公司", description="", audit_period="2024年度",
                              vouchers=list(vouchers), bank_flows=list(flows), contracts=list(contracts))


def test_revenue_net_excludes_tax_and_cost():
    v = _v("V1", "2024-03-01", [("1122", "应收账款-甲", 113, 0), ("6001", "主营业务收入", 0, 100),
                                ("2221", "应交税费-销项", 0, 13), ("6401", "主营业务成本", 70, 0), ("1405", "库存商品", 0, 70)])
    assert revenue_net(v) == 100


def test_book_without_bank_and_third_party_receipt():
    vs = [_v("V1", "2024-05-01", [("1002", "银行存款", 5_000_000, 0), ("1122", "应收账款-甲公司", 0, 5_000_000)]),
          _v("V2", "2024-06-01", [("1002", "银行存款", 3_000_000, 0), ("1122", "应收账款-乙公司", 0, 3_000_000)])]
    flows = [BankFlowRecord(transaction_id="B1", transaction_time="2024-06-01 10:00", counterparty_name="丙咨询有限公司",
                            counterparty_account="", amount=3_000_000, balance_after=0, remark="代乙公司付货款")]
    case = _case(vs, flows)
    bb = reconcile_book_to_bank(case)
    assert [x["voucher_id"] for x in bb["book_without_bank"]] == ["V1"]
    tp = detect_third_party_and_prepay(case)
    assert tp["third_party_receipts"][0]["transaction_id"] == "B1"


def test_cutoff_with_post_period_installation_payment():
    c = BusinessContract(contract_id="HT1", customer_or_vendor="某县政府", sign_date="2024-06-01", total_amount=10_000_000,
                         payment_terms="", delivery_condition="系统安装调试完成并经最终客户验收")
    v = _v("V9", "2024-12-28", [("1122", "应收账款-某县政府", 11_300_000, 0), ("6001", "主营业务收入", 0, 10_000_000),
                                ("2221", "应交税费-销项", 0, 1_300_000)], doc="HT1")
    v.entries[1].summary = "到货签收确认收入"
    flows = [BankFlowRecord(transaction_id="B9", transaction_time="2025-02-10 10:00", counterparty_name="施工队",
                            counterparty_account="", amount=-800_000, balance_after=0, remark="项目安装施工劳务分包款")]
    res = verify_cutoff_and_cost(_case([v], flows, [c]))
    assert res["cutoff_items"] and res["cutoff_items"][0]["risk"] == "HIGH"
    assert res["cutoff_revenue"] == 10_000_000


def test_materiality_uses_revenue_benchmark():
    case = get_case_by_id("REAL_CLEAN_027")
    m = compute_materiality(case)
    assert m["overall_materiality"] == round(case.financial_summary.revenue * 0.005, 2)
    assert m["performance_materiality"] == round(m["overall_materiality"] * 0.75, 2)


def test_calibration_downgrades_unanchored_model_finding():
    case = get_case_by_id("REAL_CLEAN_027")
    plugin = AuditFraudPlugin()
    tr = plugin.execute_tools(case)
    f = RiskFinding(finding_id="RF-X", title="主观怀疑", risk_level=RiskLevel.HIGH, accounting_standard="CSA 1141",
                    impact_amount=1.0, suspected_mechanism="", evidences=[], suggested_procedure="")
    cal = calibrate_model_findings([f], case, tr)
    assert f.risk_level == RiskLevel.LOW and cal["recalibrated"]


def test_offline_report_has_balanced_adjustments_and_no_clean_findings():
    h = AccountingAgentHarness()
    r = h.run_case(get_case_by_id("REAL_CSRC_001"))
    assert r.findings and r.adjustments.get("entries")
    for a in r.adjustments["entries"]:
        assert abs(sum(l["debit"] for l in a["lines"]) - sum(l["credit"] for l in a["lines"])) < 0.01
    clean = h.run_case(get_case_by_id("REAL_CLEAN_026"))
    assert clean.findings == [] and clean.overall_risk_rating == RiskLevel.CLEAN


def test_no_company_specific_templates_in_reasoning_code():
    names = ["康美", "康得新", "辅仁", "宜华", "金正大", "胜通", "东方金钰", "同济堂", "瑞幸", "索菱", "科融",
             "亚太药业", "抚顺特钢", "獐子岛", "浪奇", "雏鹰", "万福生科", "必康", "华泽", "凯迪", "三圣", "金亚",
             "欣泰", "雅百特", "风神", "茅台", "福耀"]
    for f in ["src/core/llm_adapter.py", "src/core/harness.py", "src/plugins/audit_fraud_plugin/plugin.py",
              "src/plugins/audit_fraud_plugin/rule_reasoner.py", "src/plugins/audit_fraud_plugin/substantive.py"]:
        code = "\n".join(l for l in Path(f).read_text(encoding="utf-8").splitlines() if not l.strip().startswith("#"))
        for n in names:
            assert n not in code, (f, n)


def test_contract_upload_slot():
    csv = "合同编号,交易对手方,签订日期,合同金额,付款条款,交付验收条款,是否关联方\nHT-9,甲公司,2025-01-02,100000,验收后付款,终验后转移控制权,是\n"
    vcsv = io.BytesIO()
    import pandas as pd
    pd.DataFrame([{"凭证号": "V-1", "记账日期": "2025-01-05", "科目编码": "1002", "科目名称": "银行存款", "借方金额": 100, "贷方金额": 0},
                  {"凭证号": "V-1", "记账日期": "2025-01-05", "科目编码": "6001", "科目名称": "主营业务收入", "借方金额": 0, "贷方金额": 100}]).to_excel(vcsv, index=False)
    resp = client.post("/api/upload", files={
        "vouchers_file": ("v.xlsx", vcsv.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "contracts_file": ("c.csv", csv.encode("utf-8-sig"), "text/csv")})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["counts"]["contracts"] == 1
    assert body["case"]["contracts"][0]["is_related_party"] is True
    Path(f"data/cases/custom_cases/{body['case_id']}.json").unlink(missing_ok=True)


def test_sanitize_persists_and_restores():
    from src.benchmark import test_cases as tc
    cid = "REAL_CSRC_014"
    tc._SANITIZED_OVERRIDES.pop(cid, None)
    original = get_case_by_id(cid).model_dump()
    resp = client.post(f"/api/compliance/sanitize-case/{cid}")
    assert resp.status_code == 200
    sanitized = resp.json()["data"]
    assert get_case_by_id(cid).model_dump() == sanitized      # 刷新后仍为脱敏态
    assert client.post(f"/api/compliance/restore-case/{cid}").status_code == 200
    assert get_case_by_id(cid).model_dump() == original


def test_topology_endpoint_is_data_driven():
    t = client.get("/api/topology/REAL_CSRC_001").json()
    assert t["company"].startswith("康美") and t["chains"]
    clean = client.get("/api/topology/REAL_CLEAN_027").json()
    assert all(p["flagged_amount"] == 0 for p in clean["parties"])


def test_export_carries_cpa_signoff():
    h = AccountingAgentHarness()
    rep = h.run_case(get_case_by_id("REAL_CSRC_001")).model_dump(mode="json")
    rep["cpa_signoff"] = {"auditor_name": "测试注会", "cert_no": "000000000001", "audit_firm": "测试所",
                          "sign_timestamp": "2026-10-16 10:00:00", "verification_hash": "ABC123",
                          "opinion_type": "保留意见", "professional_skepticism_notes": "n", "ethical_declaration": "e"}
    resp = client.post("/api/export/excel", json=rep)
    assert resp.status_code == 200
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "CSA1131_AI过程记录" in wb.sheetnames and "拟定审计调整分录" in wb.sheetnames
    ws = wb["CPA终审签章与合规声明"]
    assert any("测试注会" == (c.value or "") for row in ws.iter_rows() for c in row)
